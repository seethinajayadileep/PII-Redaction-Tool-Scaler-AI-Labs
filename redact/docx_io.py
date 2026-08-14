"""Read and redact .docx files in place, keeping layout as much as possible."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
from docx.text.run import Run

from redact.apply import (
    ReplacementTable,
    apply_redactions,
    replace_span_in_indexed_runs,
    resolve_overlaps,
)
from redact.detectors import detect_all
from redact.ocr_images import (
    image_is_sensitive,
    ocr_image_bytes,
    placeholder_image,
    require_ocr,
)

W_P = qn("w:p")
W_R = qn("w:r")
W_T = qn("w:t")
W_TAB = qn("w:tab")
W_BR = qn("w:br")
W_CR = qn("w:cr")
A_BLIP = qn("a:blip")
R_EMBED = qn("r:embed")


def _run_has_text(node) -> bool:
    return (
        node.find(W_T) is not None
        or node.find(W_TAB) is not None
        or node.find(W_BR) is not None
        or node.find(W_CR) is not None
    )


def _walk_text_runs(element, acc: list) -> None:
    for child in element:
        if child.tag == W_P and child is not element:
            continue
        if child.tag == W_R:
            if _run_has_text(child):
                acc.append(child)
            continue
        if child.tag == qn("w:drawing"):
            continue
        _walk_text_runs(child, acc)


def _paragraph_runs(paragraph: Paragraph) -> list[Run]:
    nodes: list = []
    _walk_text_runs(paragraph._element, nodes)
    return [Run(node, paragraph) for node in nodes]


def _story_segments(doc: Document) -> list[tuple[object | None, str]]:
    """Runs in document order, with a newline between paragraphs.

    Word often wraps a company name across two paragraphs (for example
    ``EVEREST`` then ``FAMILY TRUST``). The extra newline is not written back;
    it only exists so detectors can see the full name.
    """
    segments: list[tuple[object | None, str]] = []
    for paragraph in iter_docx_paragraphs(doc):
        if segments:
            segments.append((None, "\n"))
        for run in _paragraph_runs(paragraph):
            segments.append((run, run.text or ""))
    return segments


def _index_segments(segments: list[tuple[object | None, str]]):
    pieces: list[str] = []
    index: list[tuple[object | None, int | None]] = []
    for run, text in segments:
        pieces.append(text)
        for offset, _ch in enumerate(text):
            index.append((run, offset if run is not None else None))
    return "".join(pieces), index


def _redact_all_text(doc: Document, table: ReplacementTable) -> dict[str, int]:
    """Detect PII on the full story text and patch only the matching run characters."""
    text, index = _index_segments(_story_segments(doc))
    if not text.strip():
        return {}
    spans = resolve_overlaps(detect_all(text))
    counts: dict[str, int] = {}
    for span in sorted(spans, key=lambda item: item["start"], reverse=True):
        replacement = table.replace(span["text"], span["type"])
        replace_span_in_indexed_runs(index, span["start"], span["end"], replacement)
        counts[span["type"]] = counts.get(span["type"], 0) + 1
    return counts


def _iter_paragraphs(root, parent) -> list[Paragraph]:
    seen: set[int] = set()
    found: list[Paragraph] = []
    for node in root.iter(W_P):
        ident = id(node)
        if ident in seen:
            continue
        seen.add(ident)
        found.append(Paragraph(node, parent))
    return found


def _story_parts(doc: Document):
    """Document body plus headers, footers, notes, and comments. Each part once."""
    seen: set[int] = set()
    queue = [doc.part]
    while queue:
        part = queue.pop()
        ident = id(part)
        if ident in seen:
            continue
        seen.add(ident)
        if getattr(part, "element", None) is not None:
            yield part
        rels = getattr(part, "rels", None)
        if rels is None:
            continue
        for rel in rels.values():
            if rel.is_external:
                continue
            reltype = rel.reltype.lower()
            if any(
                name in reltype
                for name in ("header", "footer", "footnotes", "endnotes", "comments")
            ):
                queue.append(rel.target_part)


def iter_docx_paragraphs(doc: Document):
    """Body, tables, nested tables, headers, footers, and text boxes. Each w:p once."""
    seen: set[int] = set()
    for part in _story_parts(doc):
        parent = doc if part is doc.part else part
        for paragraph in _iter_paragraphs(part.element, parent):
            ident = id(paragraph._element)
            if ident in seen:
                continue
            seen.add(ident)
            yield paragraph


def extract_docx_text(path: str | Path) -> str:
    doc = Document(str(path))
    return "\n".join(p.text for p in iter_docx_paragraphs(doc) if p.text.strip())


def _image_format(content_type: str, partname: str) -> str:
    lowered = f"{content_type} {partname}".lower()
    if "jpeg" in lowered or "jpg" in lowered:
        return "JPEG"
    return "PNG"


def _redact_images(doc: Document) -> int:
    """Replace identity-document pictures with a placeholder. Logos are kept."""
    removed = 0
    scanned: dict[str, bool] = {}
    placeholders = {"PNG": placeholder_image("PNG"), "JPEG": placeholder_image("JPEG")}

    for part in _story_parts(doc):
        for blip in part.element.iter(A_BLIP):
            embed = blip.get(R_EMBED)
            if not embed or embed not in part.rels:
                continue
            image_part = part.rels[embed].target_part
            key = str(getattr(image_part, "partname", id(image_part)))
            if key not in scanned:
                blob = image_part.blob
                ocr_text = ocr_image_bytes(blob)
                scanned[key] = image_is_sensitive(ocr_text)
                if scanned[key]:
                    fmt = _image_format(getattr(image_part, "content_type", ""), key)
                    image_part._blob = placeholders[fmt]
                    image_part._image = None
                    removed += 1
            if scanned[key]:
                _set_placeholder_alt(blip)
    return removed


def _set_placeholder_alt(blip) -> None:
    """Replace drawing alt text so original ID details are not left in XML."""
    node = blip
    drawing = None
    while node is not None:
        local = node.tag.split("}")[-1]
        if local in {"inline", "anchor", "drawing"}:
            drawing = node
            if local == "drawing":
                break
        node = node.getparent()
    if drawing is None:
        return
    for child in drawing.iter():
        local = child.tag.split("}")[-1]
        if local not in {"docPr", "cNvPr"}:
            continue
        child.set("descr", "SENSITIVE IMAGE REMOVED")
        if child.get("name"):
            child.set("name", "Sensitive image removed")
        if child.get("title") is not None:
            child.set("title", "SENSITIVE IMAGE REMOVED")


def _redact_hidden_strings(doc: Document, table: ReplacementTable) -> None:
    """Replace PII in field codes, mailto links, and drawing alt text."""
    for part in _story_parts(doc):
        for node in part.element.iter(qn("w:instrText")):
            text = node.text or ""
            if not text.strip():
                continue
            new_text, _, table = apply_redactions(text, detect_all(text), table)
            if new_text != text:
                node.text = new_text
        for node in part.element.iter():
            local = node.tag.split("}")[-1]
            if local not in {"docPr", "cNvPr"}:
                continue
            for attr in ("descr", "title"):
                value = node.get(attr)
                if not value:
                    continue
                new_value, _, table = apply_redactions(value, detect_all(value), table)
                if new_value != value:
                    node.set(attr, new_value)
        for rel in part.rels.values():
            if not rel.is_external:
                continue
            target = rel.target_ref
            if not target:
                continue
            new_target, _, table = apply_redactions(target, detect_all(target), table)
            if new_target != target:
                rel._target = new_target


def redact_docx(
    source: str | Path,
    destination: str | Path,
    table: ReplacementTable | None = None,
) -> tuple[list[str], dict[str, int], ReplacementTable, dict]:
    """Copy the DOCX, replace PII in text and sensitive images, save."""
    source = Path(source)
    destination = Path(destination)
    destination.parent.mkdir(parents=True, exist_ok=True)

    original_text = extract_docx_text(source)
    doc = Document(str(source))
    table = table or ReplacementTable()
    counts = _redact_all_text(doc, table)

    _redact_hidden_strings(doc, table)

    image_count = sum(1 for part in _story_parts(doc) for _ in part.element.iter(A_BLIP))
    images_removed = 0
    if image_count:
        require_ocr()
        images_removed = _redact_images(doc)

    doc.save(str(destination))
    extra = {"images_removed": images_removed, "ocr_used": bool(image_count)}
    return [original_text], counts, table, extra
